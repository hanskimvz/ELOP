"""Result files: CSV and Excel (doc/04-software-design.md 5절)."""

import csv
import datetime
import math
import os
import platform
from dataclasses import fields
from typing import Dict, List, Optional

from . import __version__
from .config import SETUPS, TestConfig
from .error_analysis import ByteError, analyze_records, format_report
from .latency_test import Record
from .serial_port import find_port
from .stats import Stats

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS_DIR = os.path.join(PROJECT_ROOT, "results")

RECORD_FIELDS = [f.name for f in fields(Record)]
ERROR_FIELDS = ["baudrate", "seq", "kind", "pos", "tx", "rx", "xor", "tx_bin", "rx_bin",
                "subtype", "flipped_bits", "zero_to_one", "one_to_zero", "edge_adjacent"]


def _hex(v) -> str:
    return "" if v is None else "0x{:02X}".format(v)


def _bin(v) -> str:
    return "" if v is None else "{:08b}".format(v)


def error_row(e: ByteError) -> list:
    return [e.baudrate, e.seq, e.kind, e.pos, _hex(e.tx), _hex(e.rx), _hex(e.xor),
            _bin(e.tx), _bin(e.rx), e.subtype, " ".join(str(b) for b in e.flipped_bits),
            e.zero_to_one, e.one_to_zero, e.edge_adjacent]


def save_errors_csv(path: str, errors: List[ByteError]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(ERROR_FIELDS)
        for e in errors:
            w.writerow(error_row(e))


def build_metadata(cfg: TestConfig, started: datetime.datetime) -> Dict[str, str]:
    port = find_port(cfg.port)
    meta = {
        "program_version": __version__,
        "started": started.isoformat(timespec="seconds"),
        "pc": platform.node(),
        "port": cfg.port,
        "adapter": "{} ({})".format(port.adapter, port.description) if port else "unknown",
        "setup": "{} - {}".format(cfg.setup, SETUPS.get(cfg.setup, "")),
        "channel_no": str(cfg.channel_no),
        "fiber_length_m": str(cfg.fiber_length_m),
        "uart": "{}{}{} rtscts={}".format(cfg.bytesize, cfg.parity, cfg.stopbits, cfg.rtscts),
        "payload": "{} byte, pattern={}{}, seq={}".format(
            cfg.payload_size, cfg.pattern,
            " 0x{:02X}".format(cfg.fixed_byte) if cfg.pattern == "fixed" else "", cfg.use_seq),
        "iterations": str(cfg.iterations),
        "warmup": str(cfg.warmup),
        "interval_ms": str(cfg.interval_ms),
        "timeout_ms": str(cfg.timeout_ms),
        "read_mode": cfg.read_mode,
        "baudrates": ",".join(str(b) for b in cfg.baudrates()),
        "note": cfg.note,
    }
    return meta


def default_basename(cfg: TestConfig, started: datetime.datetime) -> str:
    return "latency_{}_ch{}_{}".format(cfg.setup, cfg.channel_no, started.strftime("%Y%m%d_%H%M%S"))


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return "" if math.isnan(v) else "{:.3f}".format(v)
    return str(v)


def save_csv(path: str, meta: Dict[str, str], records: List[Record]) -> None:
    """Metadata as '# key: value' lines, then one row per record."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        for k, v in meta.items():
            f.write("# {}: {}\n".format(k, v))
        w = csv.writer(f)
        w.writerow(RECORD_FIELDS)
        for r in records:
            w.writerow([_fmt(getattr(r, name)) for name in RECORD_FIELDS])


def save_xlsx(path: str, meta: Dict[str, str], records: List[Record], stats: List[Stats],
              errors: Optional[List[ByteError]] = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metadata"])
    ws["A1"].font = bold
    for k, v in meta.items():
        ws.append([k, v])
    ws.append([])

    if stats:
        ws.append(["Statistics (OK records only)"])
        ws.cell(row=ws.max_row, column=1).font = bold
        header = ["item", "unit"] + [str(s.baudrate) for s in stats]
        ws.append(header)
        for c in range(1, len(header) + 1):
            ws.cell(row=ws.max_row, column=c).font = bold
        rows_per_stat = [s.as_rows() for s in stats]
        for i, (label, _, unit) in enumerate(rows_per_stat[0]):
            values = []
            for rows in rows_per_stat:
                v = rows[i][1]
                values.append(None if isinstance(v, float) and math.isnan(v) else v)
            ws.append([label, unit] + values)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40

    raw = wb.create_sheet("Raw")
    raw.append(RECORD_FIELDS)
    for c in range(1, len(RECORD_FIELDS) + 1):
        raw.cell(row=1, column=c).font = bold
    for r in records:
        raw.append([getattr(r, name) for name in RECORD_FIELDS])
    raw.freeze_panes = "A2"

    if errors:
        es = wb.create_sheet("Errors")
        es.append(ERROR_FIELDS)
        for c in range(1, len(ERROR_FIELDS) + 1):
            es.cell(row=1, column=c).font = bold
        for e in errors:
            es.append(error_row(e))
        es.freeze_panes = "A2"

        ea = wb.create_sheet("ErrorAnalysis")
        for line in format_report(records, errors).splitlines():
            ea.append([line])
        ea.column_dimensions["A"].width = 100

    wb.save(path)


def save_results(cfg: TestConfig, started: datetime.datetime, records: List[Record],
                 stats: List[Stats], out_dir: Optional[str] = None,
                 xlsx: bool = True) -> List[str]:
    out_dir = out_dir or RESULTS_DIR
    base = os.path.join(out_dir, default_basename(cfg, started))
    meta = build_metadata(cfg, started)
    errors = analyze_records(records)
    paths = [base + ".csv"]
    save_csv(paths[0], meta, records)
    if errors:
        paths.append(base + "_errors.csv")
        save_errors_csv(paths[-1], errors)
    if xlsx:
        paths.append(base + ".xlsx")
        save_xlsx(paths[-1], meta, records, stats, errors)
    return paths


def load_csv(path: str):
    """Returns (meta, rows) where rows are dicts with numeric fields converted."""
    meta: Dict[str, str] = {}
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        lines = []
        for line in f:
            if line.startswith("#"):
                key, _, value = line[1:].partition(":")
                meta[key.strip()] = value.strip()
            else:
                lines.append(line)
    for row in csv.DictReader(lines):
        for key in ("channel", "baudrate", "seq", "rx_len", "err_bytes", "stale_bytes", "comm_flags"):
            row[key] = int(row[key]) if row.get(key) else 0
        for key in ("first_byte_us", "latency_us"):
            row[key] = float(row[key]) if row.get(key) else None
        rows.append(row)
    return meta, rows
